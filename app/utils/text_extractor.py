"""
Extracts raw text from supported file formats.
Supports: .pdf, .docx, .txt, .xlsx (basic)
"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()

    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        elif ext == ".docx":
            return _extract_docx(file_path)
        elif ext == ".txt":
            return _extract_txt(file_path)
        elif ext == ".xlsx":
            return _extract_xlsx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    except Exception as e:
        logger.error(f"Text extraction failed for {file_path}: {e}")
        raise


def _extract_pdf(path: str) -> str:
    import PyPDF2

    text_parts = []
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text.strip())

    return "\n\n".join(text_parts)


def _extract_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_xlsx(path: str) -> str:
    """
    Converts each sheet to a tabular text representation.
    Simple but covers most financial spreadsheets.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            row_str = "\t".join(str(cell) if cell is not None else "" for cell in row)
            if row_str.strip():
                parts.append(row_str)
    return "\n".join(parts)
